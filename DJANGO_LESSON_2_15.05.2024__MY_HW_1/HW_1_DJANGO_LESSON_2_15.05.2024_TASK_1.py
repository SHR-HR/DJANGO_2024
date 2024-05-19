# ВАРИАНТ - №2. (Источник для поска и вдохновения, ХабрХабр, GitHub, Reddit и поиск гугл.)

import requests
import json
import pandas as pd
from time import sleep

# Убедитесь, что openpyxl установлен
try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Use pip install openpyxl to install.")

def get_vacancies(query, area_id='40'):
    url = 'https://api.hh.ru/vacancies'
    params = {
        'text': query,
        'area': area_id,
        'per_page': 20,
        'page': 0
    }
    response = requests.get(url, params=params)
    vacancies = response.json()
    return vacancies['items']

def process_and_save(vacancies, filename):
    results = []
    for vacancy in vacancies:
        salary_from = vacancy['salary']['from'] if 'salary' in vacancy and vacancy['salary'] else None
        salary_to = vacancy['salary']['to'] if 'salary' in vacancy and vacancy['salary'] else None
        results.append([
            vacancy['id'],
            vacancy.get('name', 'No Name'),
            vacancy.get('employer', {}).get('name', 'No Employer'),
            salary_from,
            salary_to,
            vacancy.get('snippet', {}).get('requirement', 'No Details'),
            vacancy.get('snippet', {}).get('responsibility', 'No Details')
        ])

    df = pd.DataFrame(results, columns=[
        'ID', 'Name', 'Employer', 'Salary From', 'Salary To', 'Requirements', 'Responsibilities'
    ])
    df.to_excel(filename, index=False)

categories = {
    'Junior Django Developer': 'Junior Django Developer',
    'Middle Django Developer': 'Middle Django Developer',
    'Senior Django Developer': 'Senior Django Developer'
}

for title, query in categories.items():
    print(f"Processing {title}")
    vacancies = get_vacancies(query)
    if vacancies:
        process_and_save(vacancies, f'{title.replace(" ", "_")}_Jobs.xlsx')
    sleep(1)  # to avoid rate limits





# ВАРИАНТ - №1.

# import requests
# from bs4 import BeautifulSoup
# from docx import Document
#
# def fetch_job_links(query):
#     base_url = f"https://career.habr.com/vacancies?q={query.replace(' ', '%20')}&type=all"
#     response = requests.get(base_url)
#     soup = BeautifulSoup(response.text, 'html.parser')
#     links = [a['href'] for a in soup.find_all('a', href=True) if 'vacancies' in a['href']]
#     # Фильтрация, чтобы получить уникальные и полные ссылки на вакансии
#     full_links = {"https://career.habr.com" + link for link in links if link.startswith('/vacancies')}
#     return list(full_links)
#
# def fetch_job_details(url):
#     response = requests.get(url)
#     soup = BeautifulSoup(response.text, 'html.parser')
#     expectations = soup.find('div', class_='style-ugc')  # Подберите правильный селектор для ваших данных
#     return {
#         'title': soup.find('h1').text.strip() if soup.find('h1') else "No title",
#         'expectations': expectations.text.strip() if expectations else "No specific expectations listed"
#     }
#
# def create_word_document(jobs, filename):
#     doc = Document()
#     for job in jobs:
#         doc.add_heading(job['title'], level=1)
#         doc.add_paragraph(job['expectations'])
#     doc.save(filename + '.docx')
#
# # Категории разработчиков
# categories = ['Junior Django Developer', 'Middle Django Developer', 'Senior Django Developer']
# for category in categories:
#     print(f"Processing: {category}")
#     job_links = fetch_job_links(category)
#     job_details = [fetch_job_details(url) for url in job_links]
#     create_word_document(job_details, category.replace(' ', '_') + '_Jobs')